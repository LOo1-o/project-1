"""Всё, что специалисту нужно проверить, — в одном файле.

Раньше после прогона в output/ лежало больше десятка отчётов, понятных
разработчику, а важное (например, «в списке таблиц указан файл, которого
нет в папке») было только в консоли среди сотен строк. Этот модуль собирает
для человека, который готовит бюллетень:

- output/ЧТО_ПРОВЕРИТЬ.xlsx — сводка и три листа: «перед запуском»
  (файлы и список таблиц), «пустые ячейки» (где в бюллетене нет чисел и
  почему), «названия Word и Excel» (решения программы, которые стоит
  проверить глазами);
- output/Бюллетень_ПРОВЕРКА.docx — копия готового бюллетеня, где пустые по
  нашей вине ячейки подсвечены цветом. Готовый документ остаётся чистым.
"""
import re
from collections import OrderedDict
from difflib import SequenceMatcher, get_close_matches
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from openpyxl.styles import Alignment

from config import (
    _is_table_title, _normalize_text, _read_csv_rows_robustly, _table_numbers, _title_number,
    get_cleaned_cell_text, looks_like_data_value,
)
from data_filler_v2 import UNFILLED_DASH_IN_EXCEL
from logic import _one_letter_diff, _squish_text
from table_manager import TableManager

IMPORTANT = 'Важно'
FYI = 'Для сведения'
YELLOW = 'FFF2A8'   # тег есть, но числа нет по нашей вине
ORANGE = 'FFC58A'   # программа не поставила сюда тег вовсе
_YEAR_RE = re.compile(r'(19|20)\d\d')


# ---------------------------------------------------------------------------
# Названия таблиц в Word
# ---------------------------------------------------------------------------
def word_table_titles(doc) -> list:
    """Заголовки таблиц бюллетеня в порядке следования: [(номер, текст)].

    Заголовок бывает разбит на два абзаца: «5. ОБОРАЧИВАЕМОСТЬ ОБОРОТНЫХ
    АКТИВОВ ОРГАНИЗАЦИЙ» и следующим — «ПО ВИДАМ ЭКОНОМИЧЕСКОЙ
    ДЕЯТЕЛЬНОСТИ» (так в бюллетенях №1 и №2). Идущие сразу за номером
    абзацы заглавными буквами без своего номера — продолжение заголовка.
    """
    titles = []
    current = None
    for element in doc.element.body.iterchildren():
        if element.tag != qn('w:p'):
            current = None
            continue
        text = ' '.join(Paragraph(element, doc).text.split())
        if _is_table_title(text):
            current = [_title_number(text), text]
            titles.append(current)
        elif current is not None and text and text == text.upper() and any(ch.isalpha() for ch in text):
            current[1] = f"{current[1]} {text}"
        else:
            current = None
    return [tuple(title) for title in titles]


# ---------------------------------------------------------------------------
# 1. Проверка перед запуском
# ---------------------------------------------------------------------------
def _read_table_mapping(path) -> list:
    rows = [row for row in _read_csv_rows_robustly(path, delimiter=';') if row]
    if rows and str(rows[0][0]).strip().lower() in ('таблица', 'название таблицы'):
        rows = rows[1:]
    result = []
    for row in rows:
        title = str(row[0]).strip().strip('"').strip()
        source = str(row[1]).strip().strip('"').strip() if len(row) > 1 else ''
        if title:
            result.append((' '.join(title.split()), source))
    return result


def preflight(input_word_path, table_mapping_path, excel_dir, mo_path=None) -> list:
    """Проверки, которые дешевле пройти до заполнения: есть ли Word, все ли
    файлы из списка таблиц лежат в папке, совпадают ли названия таблиц со
    словами в Word. Возвращает список замечаний (пустой — всё в порядке)."""
    issues = []

    def add(level, what, where, todo):
        issues.append(OrderedDict([('Важность', level), ('Что не так', what),
                                   ('Где', where), ('Что сделать', todo)]))

    input_word_path = Path(input_word_path)
    if not input_word_path.exists():
        add(IMPORTANT, 'Нет Word-файла бюллетеня', f'папка {input_word_path.parent}',
            f'Положите Word-файл бюллетеня в папку input и назовите его ровно «{input_word_path.name}».')
        return issues
    if not Path(table_mapping_path).exists():
        add(IMPORTANT, 'Нет списка таблиц', str(table_mapping_path),
            'Нужен файл, где для каждой таблицы бюллетеня указан Excel-файл с данными '
            '(пример — в любой ветке готового бюллетеня).')
        return issues

    doc = Document(input_word_path)
    word_titles = word_table_titles(doc)
    word_by_norm = {_normalize_text(text): text for _, text in word_titles}
    word_by_number = {number: text for number, text in word_titles}
    mapping = _read_table_mapping(table_mapping_path)
    available = sorted(p.name for p in Path(excel_dir).glob('*.xlsx'))
    mapping_file = Path(table_mapping_path).name

    mapped_norms = set()
    mapped_numbers = set()
    used_sources = set()
    for title, source in mapping:
        number = _title_number(title)
        mapped_norms.add(_normalize_text(title))
        mapped_numbers.add(number)
        where = f'{mapping_file}, таблица {number if number is not None else "«" + title[:60] + "»"}'

        if _normalize_text(title) not in word_by_norm:
            word_title = word_by_number.get(number)
            if word_title is None and word_titles:
                word_title = max(word_by_norm.values(),
                                 key=lambda text: SequenceMatcher(None, _normalize_text(title), _normalize_text(text)).ratio())
            # Опечатку в одну букву программа прощает (см. get_table_name) —
            # это не повод останавливаться, но поправить стоит.
            one_letter = bool(word_title) and _one_letter_diff(_squish_text(title), _squish_text(word_title))
            add(FYI if one_letter else IMPORTANT,
                f'Название таблицы {"отличается от Word одной буквой" if one_letter else "не совпадает с Word"}: '
                f'«{title}»', where,
                f'Впишите название точно как в Word: «{word_title}».' if word_title
                else 'Такой таблицы в Word не нашлось — проверьте номер и название.')

        if not source:
            add(FYI, 'У таблицы не указан Excel-файл', where,
                'Программа эту таблицу не заполняет — заполните её вручную '
                '(или укажите Excel-файл, если он есть).')
            continue
        used_sources.add(source)
        if source not in available:
            if f'{source}.xlsx' in available:
                hint = f'В имени нет «.xlsx» в конце: должно быть «{source}.xlsx».'
            else:
                close = get_close_matches(source, available, n=1, cutoff=0.6)
                hint = (f'Такого файла нет в папке input/excel. Возможно, имелся в виду «{close[0]}» — '
                        f'впишите это имя.' if close else
                        'Такого файла нет в папке input/excel — загрузите его или исправьте имя.')
            add(IMPORTANT, f'Нет Excel-файла «{source}»', where, hint)
            used_sources.add(f'{source}.xlsx')

    for number, text in word_titles:
        if _normalize_text(text) not in mapped_norms and number not in mapped_numbers:
            add(IMPORTANT, f'Таблицы нет в списке таблиц: «{text}»', mapping_file,
                'Добавьте строку «название таблицы;имя Excel-файла» — иначе таблица не заполнится.')

    for name in available:
        if name not in used_sources and not any(get_close_matches(name, [s], cutoff=0.95) for s in used_sources):
            add(FYI, f'Excel-файл «{name}» нигде не используется', 'папка input/excel',
                f'Если он нужен для какой-то таблицы — укажите его в {mapping_file}.')

    if mo_path is not None and not Path(mo_path).exists() and any('mo' in s.lower() for _, s in mapping):
        add(IMPORTANT, 'Нет справочника муниципальных образований', str(mo_path),
            'Для таблиц по муниципальным образованиям нужен файл mo.csv (код;название).')
    return issues


# ---------------------------------------------------------------------------
# 2. Пустые ячейки и копия с подсветкой
# ---------------------------------------------------------------------------
def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def _is_label_or_header_row(cells) -> bool:
    """Служебная строка («в том числе:») или строка годов (2024 | 2025)."""
    texts = []
    first = cells[0]._tc
    prev = None
    for cell in cells[1:]:
        if cell._tc is first or cell._tc is prev:
            continue
        prev = cell._tc
        text = get_cleaned_cell_text(cell).strip()
        if text:
            texts.append(text)
    if not texts:
        return True
    return all(_YEAR_RE.fullmatch(text) for text in texts)


def mark_problem_cells(input_doc_path, template_doc_path, final_doc_path, unfilled, out_path,
                       numbers_without_source=(), unknown_row_names=()) -> list:
    """Подсвечивает в копии готового бюллетеня ячейки без чисел по нашей вине
    и возвращает их список для отчёта.

    Жёлтая — тег есть, но число не нашлось (строки или показателя нет в
    данных Excel). Оранжевая — программа не поставила тег туда, где в
    исходном Word было число (не узнана колонка или строка, или таблицу
    программа не заполняет). Ячейки, где в самом Excel прочерк, не
    подсвечиваются — это не ошибка.
    """
    source = Document(input_doc_path)
    template = Document(template_doc_path)
    final = Document(final_doc_path)
    final_tables = TableManager(final).iter_tables()
    numbers = _table_numbers(final, final_tables)
    unfilled_by_cell = {(u['Таблица'], u['Строка'], u.get('Колонка')): u for u in unfilled}
    unknown_rows = {_normalize_text(name) for name in unknown_row_names}
    numbers_without_source = set(numbers_without_source)

    problems = []
    for t_idx, (src_table, tpl_table, fin_table) in enumerate(
            zip(TableManager(source).iter_tables(), TableManager(template).iter_tables(), final_tables)):
        number = numbers[t_idx]
        for r_idx, (src_row, tpl_row, fin_row) in enumerate(zip(src_table.rows, tpl_table.rows, fin_table.rows)):
            fin_cells = fin_row.cells
            if not fin_cells:
                continue
            row_name = ' '.join(get_cleaned_cell_text(fin_cells[0]).split())
            if not row_name or _is_label_or_header_row(src_row.cells):
                continue
            first_tc = fin_cells[0]._tc
            prev_tc = None
            visible_col = 1
            for c_idx, (src_cell, tpl_cell, fin_cell) in enumerate(zip(src_row.cells, tpl_row.cells, fin_cells)):
                if fin_cell._tc is first_tc or fin_cell._tc is prev_tc:
                    continue
                prev_tc = fin_cell._tc
                visible_col += 1
                entry = unfilled_by_cell.get((t_idx + 1, r_idx + 1, c_idx + 1))
                if entry is not None:
                    if entry.get('Причина') == UNFILLED_DASH_IN_EXCEL:
                        continue
                    color, why, indicator = YELLOW, entry.get('Причина', ''), entry.get('Показатель', '')
                elif '{{' not in tpl_cell.text and looks_like_data_value(get_cleaned_cell_text(src_cell)):
                    indicator = ''
                    color = ORANGE
                    if number in numbers_without_source or number is None:
                        why = 'Эту таблицу программа не заполняет (в списке таблиц нет Excel-файла) — заполните вручную'
                    elif _normalize_text(row_name) in unknown_rows:
                        why = 'Название строки не узнано — см. лист «3 Названия Word и Excel»'
                    else:
                        why = ('Колонка не узнана: её название в Word не совпало ни с одним показателем Excel — '
                               'см. лист «3 Названия Word и Excel» и раздел «Частые случаи» инструкции')
                else:
                    continue
                _shade(fin_cell, color)
                problems.append({'Таблица': number, 'Строка в Word': row_name, 'Колонка': visible_col,
                                 'Показатель': indicator, 'Почему пусто': why,
                                 'Цвет в Бюллетень_ПРОВЕРКА.docx': 'жёлтый' if color == YELLOW else 'оранжевый'})
    final.save(out_path)
    return problems


def _group_problems(problems) -> list:
    groups = OrderedDict()
    for p in problems:
        key = (p['Таблица'], p['Строка в Word'], p['Почему пусто'], p['Цвет в Бюллетень_ПРОВЕРКА.docx'])
        group = groups.setdefault(key, {'cols': [], 'indicators': []})
        group['cols'].append(p['Колонка'])
        if p['Показатель'] and p['Показатель'] not in group['indicators']:
            group['indicators'].append(p['Показатель'])
    rows = []
    for (number, row_name, why, color), group in groups.items():
        cols = sorted(set(group['cols']))
        rows.append(OrderedDict([
            ('Таблица', number if number is not None else '—'),
            ('Строка в Word', row_name),
            ('Колонки', ', '.join(map(str, cols))),
            ('Сколько ячеек', len(group['cols'])),
            ('Показатель', '; '.join(group['indicators'])),
            ('Почему пусто', why),
            ('Цвет в Бюллетень_ПРОВЕРКА.docx', color),
        ]))
    return rows


# ---------------------------------------------------------------------------
# 3. Сборка ЧТО_ПРОВЕРИТЬ.xlsx
# ---------------------------------------------------------------------------
def write_check_workbook(path, preflight_issues, problems, name_check_report_path=None,
                         check_doc_name='Бюллетень_ПРОВЕРКА.docx', final_doc_name='Бюллетень_ГОТОВЫЙ.docx'):
    path = Path(path)
    grouped = _group_problems(problems)
    names = pd.DataFrame()
    if name_check_report_path and Path(name_check_report_path).exists():
        names = pd.read_excel(name_check_report_path)

    important = sum(1 for i in preflight_issues if i['Важность'] == IMPORTANT)
    yellow = sum(1 for p in problems if p['Цвет в Бюллетень_ПРОВЕРКА.docx'] == 'жёлтый')
    orange = len(problems) - yellow
    summary = pd.DataFrame([
        ('1 Перед запуском', f'{important} важных замечаний, {len(preflight_issues) - important} для сведения',
         'Ошибки в списке таблиц и файлах. Исправьте важные и запустите программу снова.'),
        ('2 Пустые ячейки', f'{len(problems)} ячеек: {orange} оранжевых, {yellow} жёлтых',
         f'Где в бюллетене нет чисел по нашей вине и почему. Те же ячейки подсвечены в {check_doc_name}. '
         f'Прочерки, которые есть в самом Excel, сюда не попадают — они верные.'),
        ('3 Названия Word и Excel', f'{len(names)} строк',
         'Где программа сама решила, что названия в Word и Excel означают одно и то же, и строки, '
         'которые она не узнала. Проверьте глазами; что поправить — написано в каждой строке.'),
        ('Готовый бюллетень', final_doc_name,
         f'Чистый документ без подсветки — его и отправляйте. {check_doc_name} — только для проверки.'),
    ], columns=['Лист', 'Сколько', 'Что это'])

    sheets = OrderedDict([
        ('Сводка', summary),
        ('1 Перед запуском', pd.DataFrame(preflight_issues) if preflight_issues
         else pd.DataFrame([{'Что не так': 'Замечаний нет'}])),
        ('2 Пустые ячейки', pd.DataFrame(grouped) if grouped
         else pd.DataFrame([{'Почему пусто': 'Пустых по нашей вине ячеек нет'}])),
        ('3 Названия Word и Excel', names if not names.empty
         else pd.DataFrame([{'Что случилось': 'Замечаний нет'}])),
    ])
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            for column_cells in sheet.columns:
                header = str(column_cells[0].value or '')
                longest = max(len(str(c.value or '')) for c in column_cells)
                width = 60 if longest > 60 else max(12, min(longest, 60), len(header) + 2)
                sheet.column_dimensions[column_cells[0].column_letter].width = width
                for cell in column_cells:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
    return {'important': important, 'problems': len(problems), 'orange': orange, 'yellow': yellow,
            'names': len(names)}
