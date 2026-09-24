"""Проверка перед запуском и подсветка пустых ячеек (check_report.py).

Случаи из бюллетеня №3, как его загрузили: у файла нет «.xlsx», указан
файл прошлого бюллетеня (M25 вместо S25), название таблицы не совпадает с
Word. Программа замечала только первые два и лишь в консоли.
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

from check_report import IMPORTANT, mark_problem_cells, preflight, word_table_titles, write_check_workbook
from data_filler_v2 import UNFILLED_DASH_IN_EXCEL, UNFILLED_NO_INDICATOR


def _word(path, titles):
    doc = Document()
    for title in titles:
        for part in title:
            doc.add_paragraph(part)
        doc.add_table(rows=1, cols=2)
    doc.save(path)


class PreflightTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        (self.dir / 'excel').mkdir()
        for name in ('S25_t22mo8.xlsx', 'S25_t27.xlsx', 'S25_t23.xlsx', 'S25_лишний.xlsx'):
            openpyxl.Workbook().save(self.dir / 'excel' / name)
        self.word = self.dir / 'Бюллетень.docx'
        _word(self.word, [
            ['1. ОСНОВНЫЕ ПОКАЗАТЕЛИ'],
            ['13. ФИНАНСОВАЯ УСТОЙЧИВОСТЬ'],
            ['14. ФОРМИРОВАНИЕ РЕЗУЛЬТАТОВ ОРГАНИЗАЦИЙ', 'ПО ВИДАМ ЭКОНОМИЧЕСКОЙ ДЕЯТЕЛЬНОСТИ'],
            ['20. РЕНТАБЕЛЬНОСТЬ АКТИВОВ'],
        ])
        self.mapping = self.dir / 'table_source_data_mapping.csv'

    def tearDown(self):
        self.tmp.cleanup()

    def _issues(self, lines):
        self.mapping.write_text('Таблица;Источник\n' + '\n'.join(lines), encoding='utf-8')
        return preflight(self.word, self.mapping, self.dir / 'excel')

    def test_split_title_is_joined(self):
        titles = dict(word_table_titles(Document(self.word)))
        self.assertEqual(titles[14], '14. ФОРМИРОВАНИЕ РЕЗУЛЬТАТОВ ОРГАНИЗАЦИЙ ПО ВИДАМ ЭКОНОМИЧЕСКОЙ ДЕЯТЕЛЬНОСТИ')

    def test_real_mistakes_are_found_with_hints(self):
        issues = self._issues([
            '"1. ОСНОВНЫЕ ПОКАЗАТЕЛИ";',
            '"13. ФИНАНСОВАЯ УСТОЙЧИВОСТЬ";S25_t22mo8',
            '"14. ФОРМИРОВАНИЕ РЕЗУЛЬТАТОВ ОРГАНИЗАЦИЙ ПО ОРГАНИЗАЦИОННО-ПРАВОВЫМ ФОРМАМ";S25_t23.xlsx',
            '"20. РЕНТАБЕЛЬНОСТЬ АКТИВОВ";M25_t27.xlsx',
        ])
        text = '\n'.join(f"{i['Важность']}|{i['Что не так']}|{i['Что сделать']}" for i in issues)
        self.assertIn('должно быть «S25_t22mo8.xlsx»', text)
        self.assertIn('Возможно, имелся в виду «S25_t27.xlsx»', text)
        self.assertIn('«14. ФОРМИРОВАНИЕ РЕЗУЛЬТАТОВ ОРГАНИЗАЦИЙ ПО ВИДАМ ЭКОНОМИЧЕСКОЙ ДЕЯТЕЛЬНОСТИ»', text)
        self.assertIn('Программа эту таблицу не заполняет', text)
        self.assertIn('«S25_лишний.xlsx» нигде не используется', text)
        self.assertEqual(sum(i['Важность'] == IMPORTANT for i in issues), 3)

    def test_clean_input_has_no_important_issues(self):
        issues = self._issues([
            '"1. ОСНОВНЫЕ ПОКАЗАТЕЛИ";S25_лишний.xlsx',
            '"13. ФИНАНСОВАЯ УСТОЙЧИВОСТЬ";S25_t22mo8.xlsx',
            '"14. ФОРМИРОВАНИЕ РЕЗУЛЬТАТОВ ОРГАНИЗАЦИЙ ПО ВИДАМ ЭКОНОМИЧЕСКОЙ ДЕЯТЕЛЬНОСТИ";S25_t23.xlsx',
            '"20. РЕНТАБЕЛЬНОСТЬ АКТИВОВ";S25_t27.xlsx',
        ])
        self.assertEqual([i for i in issues if i['Важность'] == IMPORTANT], [])

    def test_missing_word(self):
        self.word.unlink()
        issues = self._issues(['"1. ОСНОВНЫЕ ПОКАЗАТЕЛИ";S25_t27.xlsx'])
        self.assertEqual(issues[0]['Что не так'], 'Нет Word-файла бюллетеня')


class MarkProblemCellsTest(unittest.TestCase):
    def test_colors_and_reasons(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            rows = [['', '2024', '2025', 'Прочее'], ['в том числе:', '', '', ''],
                    ['Сельское хозяйство', '1', '2', '3']]
            paths = {}
            for kind, third_row in (('src', ['Сельское хозяйство', '1', '2', '3']),
                                    ('tpl', ['Сельское хозяйство', '{{A_X_24}}', '{{A_X_25}}', '']),
                                    ('fin', ['Сельское хозяйство', '-', '-', ''])):
                doc = Document()
                doc.add_paragraph('4. БАЛАНС')
                table = doc.add_table(rows=3, cols=4)
                for row, values in zip(table.rows, rows[:2] + [third_row]):
                    for cell, value in zip(row.cells, values):
                        cell.text = value
                paths[kind] = tmp / f'{kind}.docx'
                doc.save(paths[kind])
            unfilled = [
                {'Таблица': 1, 'Строка': 3, 'Колонка': 2, 'Показатель': 'X', 'Причина': UNFILLED_DASH_IN_EXCEL},
                {'Таблица': 1, 'Строка': 3, 'Колонка': 3, 'Показатель': 'X', 'Причина': UNFILLED_NO_INDICATOR},
            ]
            out = tmp / 'check.docx'
            problems = mark_problem_cells(paths['src'], paths['tpl'], paths['fin'], unfilled, out)
            by_col = {p['Колонка']: p for p in problems}
            self.assertNotIn(2, by_col, 'прочерк, который есть в самом Excel, — не проблема')
            self.assertEqual(by_col[3]['Цвет в Бюллетень_ПРОВЕРКА.docx'], 'жёлтый')
            self.assertEqual(by_col[4]['Цвет в Бюллетень_ПРОВЕРКА.docx'], 'оранжевый')
            self.assertEqual(len(problems), 2, 'строки годов и «в том числе:» не проверяются')
            self.assertEqual(problems[0]['Таблица'], 4)
            fill = Document(out).tables[0].rows[2].cells[3]._tc.tcPr.xpath('./w:shd/@w:fill')
            self.assertEqual(fill, ['FFC58A'])
            with contextlib.redirect_stdout(io.StringIO()):
                summary = write_check_workbook(tmp / 'ЧТО_ПРОВЕРИТЬ.xlsx', [], problems)
            self.assertEqual((summary['orange'], summary['yellow']), (1, 1))
            self.assertEqual(openpyxl.load_workbook(tmp / 'ЧТО_ПРОВЕРИТЬ.xlsx').sheetnames,
                             ['Сводка', '1 Перед запуском', '2 Пустые ячейки', '3 Названия Word и Excel'])


if __name__ == '__main__':
    unittest.main()
