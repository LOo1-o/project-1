import tempfile
import unittest
from pathlib import Path

from config import get_table_name
from docx import Document


class TestTableNameFuzzy(unittest.TestCase):
    def test_get_table_name_fuzzy_from_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test.docx"
            doc = Document()
            doc.add_paragraph("Раздел 1")
            doc.add_paragraph("1. Баланс организаций по видам экономической деятельности")
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = [
                "1. Баланс организаций по видам экономической деятельности",
                "2. Отчет о движении денежных средств"
            ]

            title = get_table_name(table, known_titles)
            self.assertEqual(title, known_titles[0])

    def test_get_table_name_fuzzy_with_typo(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            doc_path = Path(tmpdir) / "test_typo.docx"
            doc = Document()
            doc.add_paragraph("1. Баланс оргнизаций по видам экономики")
            doc.add_table(rows=1, cols=1)
            doc.save(doc_path)

            loaded = Document(doc_path)
            table = loaded.tables[0]
            known_titles = [
                "1. Баланс организаций по видам экономической деятельности"
            ]

            title = get_table_name(table, known_titles)
            self.assertEqual(title, known_titles[0])


if __name__ == "__main__":
    unittest.main()
