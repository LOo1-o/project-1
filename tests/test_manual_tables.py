"""Сводные таблицы без Excel-файла в списке таблиц очищаются целиком.

В бюллетене №3 у таблиц 1–3 в table_source_data_mapping.csv источник
пустой: их заполняют вручную. Программа о них не знала: таблица 1
оставалась с числами прошлого периода, таблицу 2 «узнавала» по названию
показателя «Валюта баланса» (как t01) и ставила прочерки лишь в части
строк, таблицу 3 не трогала вовсе. Теперь все числа стираются, ячейки
остаются пустыми (не прочерк — данные есть, их вносит человек), шапка
с годами и названия строк — на месте.
"""
import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document

from config import load_manual_table_numbers
from logic import generate_word_template

ROOT = Path(__file__).resolve().parent.parent
MAPPINGS = ROOT / 'input' / 'mappings'


def _table(doc, title, rows):
    doc.add_paragraph(title)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    return table


class ManualTablesTest(unittest.TestCase):
    def test_loader_takes_only_rows_without_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'table_source_data_mapping.csv'
            path.write_text('Таблица;Источник\n'
                            '"1. ОСНОВНЫЕ ФИНАНСОВЫЕ ПОКАЗАТЕЛИ";\n'
                            '"2.  БАЛАНС ОРГАНИЗАЦИЙ";  \n'
                            '"4. БАЛАНС ПО ВИДАМ ДЕЯТЕЛЬНОСТИ";S25_t01.xlsx\n'
                            'Без номера;\n', encoding='utf-8')
            self.assertEqual(load_manual_table_numbers(path),
                             {1: '1. ОСНОВНЫЕ ФИНАНСОВЫЕ ПОКАЗАТЕЛИ', 2: '2. БАЛАНС ОРГАНИЗАЦИЙ'})

    def test_numbers_cleared_to_empty_header_kept(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = Document()
            _table(doc, '1. ОСНОВНЫЕ ФИНАНСОВЫЕ ПОКАЗАТЕЛИ', [
                ['', '2025'],
                ['Выручка', ''],
                ['тысяч рублей', '609 349 963'],
                ['Сальдо прочих доходов и расходов', '-29 288 219'],
            ])
            _table(doc, '2. БАЛАНС ОРГАНИЗАЦИЙ', [
                ['', '2025', 'Справочно: 2024*'],
                ['Валюта баланса', '1 167 278 789', '912 076 882'],
                ['в том числе:', '', ''],
                ['запасы', '126 137 691', '111 902 252'],
            ])
            src = tmp / 'in.docx'
            doc.save(src)
            mapping = tmp / 'table_source_data_mapping.csv'
            mapping.write_text('Таблица;Источник\n"1. ОСНОВНЫЕ ФИНАНСОВЫЕ ПОКАЗАТЕЛИ";\n'
                               '"2. БАЛАНС ОРГАНИЗАЦИЙ";\n', encoding='utf-8')
            out = tmp / 'out.docx'
            with contextlib.redirect_stdout(io.StringIO()):
                generate_word_template(src, MAPPINGS / 'okved_mapping.csv', mapping,
                                       MAPPINGS / 'column_mapping_v2.csv', out, clear_only=True)
            texts = [[[c.text for c in row.cells] for row in t.rows] for t in Document(out).tables]
            self.assertEqual(texts[0], [['', '2025'], ['Выручка', ''], ['тысяч рублей', ''],
                                        ['Сальдо прочих доходов и расходов', '']])
            self.assertEqual(texts[1], [['', '2025', 'Справочно: 2024*'], ['Валюта баланса', '', ''],
                                        ['в том числе:', '', ''], ['запасы', '', '']],
                             'ни прочерков, ни «узнавания» таблицы по названию показателя')


if __name__ == '__main__':
    unittest.main()
